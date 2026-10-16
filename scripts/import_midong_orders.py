#!/usr/bin/env python3
"""(주)미동이엔씨 파트너 포털(partner.midong.com) 자재마감 «월별 품목 내역» 수집본 → 유니블세일 반영 + 엑셀.
입력: 수집 JSON(포털 로그인 화면에서 읽기 전용 GET으로 모은 줄들; {rows:[…], summaries:[…]}).
출력: ①data/db.json 의 midongOrders(백업 후) ②~/Desktop/동산화스너/미동발주내역/미동이엔씨_발주내역_전체_<날짜>.xlsx
사용: venv/bin/python scripts/import_midong_orders.py <raw.json> [--apply]"""
import json, sys, os, time, shutil, collections, re, html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

raw = json.load(open(sys.argv[1], encoding='utf-8'))
apply = '--apply' in sys.argv
num = lambda s: float(re.sub(r'[^\d.\-]', '', s) or 0) if isinstance(s, str) else float(s or 0)
rows = []
for r in raw['rows']:
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', r['inDate'] or ''): in_date = ''
    else: in_date = r['inDate']
    close = r['closeDate'] if re.fullmatch(r'\d{4}-\d{2}-\d{2}', r['closeDate'] or '') else ''
    rows.append({
        'id': f"mo-{r['ym']}-{int(r['no']):04d}", 'closeYm': r['ym'], 'rno': r['rno'], 'site': html.unescape(r['site']), 'name': html.unescape(r['name']), 'spec': html.unescape(r['spec']), 'unit': html.unescape(r['unit']),  # 포털 HTML 엔티티(&amp; 등) 해제
        'ordQty': num(r['ordQty']), 'price': num(r['price']), 'ordAmt': int(num(r['ordAmt'])), 'inQty': num(r['inQty']), 'inAmt': int(num(r['inAmt'])),
        'inDate': in_date, 'closeDate': close})
ids = [r['id'] for r in rows]
assert len(ids) == len(set(ids)), '중복 id'
print(f"줄 {len(rows)}개 · {min(r['closeYm'] for r in rows)}~{max(r['closeYm'] for r in rows)} · 입고금액 합계 {sum(r['inAmt'] for r in rows):,}")

# ERP 미동 매출(전표) — 달력 월별
d = json.load(open('data/db.json', encoding='utf-8'))
erp = collections.Counter(); erpn = collections.Counter()
for s in d['slips']:
    if '미동' in s['customerName']:
        sign = -1 if s['slipType'].lower().startswith('return') else 1
        erp[s['slipDate'][:7]] += sign * s['totalAmount']; erpn[s['slipDate'][:7]] += 1

wb = Workbook(); H = Font(bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='1F4E79')
def sheet(ws, head, data, widths=None, money=()):
    ws.append(head)
    for c in ws[1]: c.font = H; c.fill = FILL; c.alignment = Alignment(horizontal='center')
    for row in data: ws.append(row)
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths or [], 1): ws.column_dimensions[get_column_letter(i)].width = w
    for col in money:
        for cell in ws[get_column_letter(col)][1:]: cell.number_format = '#,##0'
ws = wb.active; ws.title = '전체 품목(입고 기준)'
sheet(ws, ['마감월', '청구번호', '현장', '품명', '규격', '단위', '발주수량', '단가', '발주금액', '입고수량', '입고금액', '입고일', '마감일'],
      [[f"{r['closeYm'][:4]}-{r['closeYm'][4:]}", r['rno'], r['site'], r['name'], r['spec'], r['unit'], r['ordQty'], r['price'], r['ordAmt'], r['inQty'], r['inAmt'], r['inDate'], r['closeDate']] for r in rows],
      [9, 12, 34, 28, 16, 6, 10, 10, 12, 10, 12, 12, 12], money=(7, 8, 9, 10, 11))
# 월별 마감 요약
by = collections.defaultdict(lambda: [0, 0, 0, 0])
for r in rows:
    b = by[r['closeYm']]; b[0] += 1; b[1] += r['ordAmt']; b[2] += r['inAmt']; b[3] += (1 if r['inQty'] != r['ordQty'] else 0)
summ = {s['ym']: s for s in raw.get('summaries', [])}
data = []
for ym in sorted(by):
    b = by[ym]; s = summ.get(ym, {})
    data.append([f"{ym[:4]}-{ym[4:]}", s.get('period', ''), b[0], b[1], b[2], b[3], s.get('inAmt', ''), s.get('inCnt', ''), s.get('paidAmt', ''), s.get('paidCnt', ''), s.get('unpaidAmt', ''), s.get('unpaidCnt', '')])
ws2 = wb.create_sheet('월별 마감 요약')
sheet(ws2, ['마감월', '기간', '줄수(수집)', '발주금액', '입고금액(수집 합계)', '발주≠입고 줄', '포털 입고금액', '포털 입고건', '포털 결제금액', '포털 결제건', '포털 미결금액', '포털 미결건'], data, [9, 16, 10, 14, 16, 12, 14, 10, 14, 10, 14, 10], money=(4, 5, 7, 9, 11))
# 청구번호별
by2 = collections.defaultdict(lambda: {'site': '', 'n': 0, 'ord': 0, 'in': 0, 'd0': '9999', 'd1': ''})
for r in rows:
    b = by2[r['rno']]; b['site'] = r['site']; b['n'] += 1; b['ord'] += r['ordAmt']; b['in'] += r['inAmt']
    if r['inDate']: b['d0'] = min(b['d0'], r['inDate']); b['d1'] = max(b['d1'], r['inDate'])
ws3 = wb.create_sheet('청구번호별')
sheet(ws3, ['청구번호', '현장', '줄수', '발주금액', '입고금액', '첫 입고일', '마지막 입고일'], [[k, v['site'], v['n'], v['ord'], v['in'], v['d0'] if v['d0'] != '9999' else '', v['d1']] for k, v in sorted(by2.items(), key=lambda x: x[1]['d1'], reverse=True)], [12, 40, 8, 14, 14, 12, 12], money=(4, 5))
# 현장별
by3 = collections.defaultdict(lambda: collections.Counter())
for r in rows: by3[r['site']][r['closeYm'][:4]] += r['inAmt']
years = sorted({y for v in by3.values() for y in v})
ws4 = wb.create_sheet('현장별 연도 입고금액')
sheet(ws4, ['현장'] + years + ['합계'], [[k] + [v[y] for y in years] + [sum(v.values())] for k, v in sorted(by3.items(), key=lambda x: -sum(x[1].values()))], [40] + [13] * (len(years) + 1), money=tuple(range(2, len(years) + 3)))
# ERP 대조(달력 월, 입고일 기준)
cal = collections.Counter(); caln = collections.Counter()
for r in rows:
    if r['inDate']: cal[r['inDate'][:7]] += r['inAmt']; caln[r['inDate'][:7]] += 1
months = sorted(set(cal) | set(erp))
ws5 = wb.create_sheet('ERP 대조(달력 월)')
sheet(ws5, ['월', '포털 입고금액(입고일 기준)', '포털 줄수', 'ERP 미동 매출(전표 일자 기준, 반품 차감)', 'ERP 전표수', '차이(포털−ERP)', '비고'],
      [[m, cal[m], caln[m], erp[m], erpn[m], cal[m] - erp[m], '' if erpn[m] else 'ERP에 해당 월 전표 없음'] for m in months], [10, 22, 10, 30, 10, 16, 24], money=(2, 4, 6))
out = os.path.expanduser(f"~/Desktop/동산화스너/미동발주내역/미동이엔씨_발주내역_전체_{time.strftime('%Y%m%d')}.xlsx")
os.makedirs(os.path.dirname(out), exist_ok=True); wb.save(out); print('엑셀', out)
if apply:
    ts = time.strftime('%Y%m%d_%H%M%S'); shutil.copy2('data/db.json', f'data/db.json.bak_midong_{ts}')
    d['midongOrders'] = rows
    json.dump(d, open('data/db.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
    json.dump({'summaries': raw.get('summaries', []), 'fetchedAt': ts}, open('data/midong_summaries.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('db.json midongOrders', len(rows), '반영(백업 data/db.json.bak_midong_%s)' % ts)
